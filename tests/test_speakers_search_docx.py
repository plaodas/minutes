import uuid
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from backend.app import app
from minutes.db import session_scope
from minutes.minutes_docx import build_minutes_docx
from minutes.models import Task, TaskHistory
from minutes.search_text import build_search_text, escape_like
from minutes.speakers import apply_speaker_updates
from minutes.task_state import update_success
from tests.auth_helpers import create_user, logged_in_client, login


def _result():
    return {
        "segments": [
            {"start": 0, "end": 2, "text": "こんにちは"},
            {"start": 2, "end": 4, "text": "了解しました"},
        ],
        "transcript": "こんにちは\n了解しました",
        "summary": "挨拶の要約",
        "minutes": "会議の本文",
        "action_items": [{"who": "山田", "what": "資料送付", "due": "金曜"}],
    }


def _add_task(user_id, *, name="定例", result=None, search_text=None, deleted=False):
    task_id = uuid.uuid4()
    payload = _result() if result is None else result
    with session_scope() as session:
        session.add(
            Task(
                id=task_id,
                user_id=user_id,
                name=name,
                status="success",
                result=payload,
                search_text=(
                    search_text
                    if search_text is not None
                    else build_search_text(name, payload)
                ),
                deleted=deleted,
            )
        )
    return task_id


def test_search_text_includes_name_transcript_summary_and_minutes():
    text = build_search_text("定例", _result())
    assert "定例" in text
    assert "こんにちは" in text
    assert "挨拶の要約" in text
    assert "会議の本文" in text
    assert escape_like("100%_done") == "100\\%\\_done"


def test_success_persists_search_text():
    _client, user_id = logged_in_client(app)
    task_id = _add_task(user_id, result={"transcript": "placeholder"})
    update_success(str(task_id), _result(), lambda *_args, **_kwargs: None)
    with session_scope() as session:
        task = session.get(Task, task_id)
        assert task is not None
        assert "挨拶の要約" in task.search_text


def test_owner_can_bulk_assign_speakers_and_history_is_recorded():
    client, user_id = logged_in_client(app)
    task_id = _add_task(user_id)
    response = client.post(
        f"/api/bg/task/{task_id}/speakers",
        json={
            "updates": [
                {"index": 0, "speaker": "山田"},
                {"index": 1, "speaker": "山田"},
            ]
        },
    )
    assert response.status_code == 200
    body = response.json()
    assert body["transcript"] == "[山田] こんにちは\n[山田] 了解しました"
    assert body["segments"][1]["speaker"] == "山田"
    with session_scope() as session:
        task = session.get(Task, task_id)
        assert "[山田]" in task.search_text
        events = session.query(TaskHistory).filter(TaskHistory.task_id == task_id).all()
    assert any(event.event_type == "speakers" for event in events)


def test_speaker_update_rejects_bad_input_and_other_owners():
    client, user_id = logged_in_client(app)
    task_id = _add_task(user_id)
    other = TestClient(app)
    other_id = create_user()
    login(other, other_id)
    assert (
        other.post(
            f"/api/bg/task/{task_id}/speakers",
            json={"updates": [{"index": 0, "speaker": "佐藤"}]},
        ).status_code
        == 404
    )
    assert (
        client.post(
            f"/api/bg/task/{task_id}/speakers",
            json={"updates": [{"index": 9, "speaker": "佐藤"}]},
        ).status_code
        == 400
    )
    assert (
        client.post(
            f"/api/bg/task/{task_id}/speakers",
            json={"updates": [{"index": 0, "speaker": "  "}]},
        ).status_code
        == 400
    )


def test_apply_speaker_updates_copies_result():
    original = _result()
    updated = apply_speaker_updates(original, [(0, "田中")])
    assert "speaker" not in original["segments"][0]
    assert updated["segments"][0]["speaker"] == "田中"


def test_search_matches_owned_fields_and_pages():
    client, user_id = logged_in_client(app)
    other_id = create_user()
    name_hit = _add_task(
        user_id, name="予算会議", result={**_result(), "summary": "別件"}
    )
    transcript_hit = _add_task(
        user_id,
        name="無関係",
        result={**_result(), "transcript": "独自キーワードの発言"},
    )
    summary_hit = _add_task(
        user_id, name="無関係", result={**_result(), "summary": "独自要約語"}
    )
    minutes_hit = _add_task(
        user_id, name="無関係", result={**_result(), "minutes": "独自議事本文"}
    )
    _add_task(other_id, name="予算会議")
    _add_task(user_id, name="予算会議", deleted=True)

    def ids_for(query, **params):
        response = client.get(
            "/api/bg/tasks", params={"q": query, "limit": 20, **params}
        )
        assert response.status_code == 200
        return {item["id"] for item in response.json()["tasks"]}

    assert str(name_hit) in ids_for("予算")
    assert str(transcript_hit) in ids_for("独自キーワード")
    assert str(summary_hit) in ids_for("独自要約")
    assert str(minutes_hit) in ids_for("独自議事")
    assert ids_for("予算") == {str(name_hit)}

    first = client.get("/api/bg/tasks", params={"q": "無関係", "limit": 1, "offset": 0})
    second = client.get(
        "/api/bg/tasks", params={"q": "無関係", "limit": 1, "offset": 1}
    )
    assert first.status_code == 200
    assert second.status_code == 200
    assert first.json()["tasks"][0]["id"] != second.json()["tasks"][0]["id"]
    empty = client.get("/api/bg/tasks", params={"q": " "})
    assert empty.status_code == 200
    short = client.get("/api/bg/tasks", params={"q": "a"})
    assert short.status_code == 400


def test_minutes_download_accepts_docx_and_keeps_txt(monkeypatch, tmp_path):
    monkeypatch.setenv("OUTPUTS_DIR", str(tmp_path))
    client, user_id = logged_in_client(app)
    output = tmp_path / "minutes.txt"
    output.write_text("会議の本文", encoding="utf-8")
    task_id = _add_task(
        user_id,
        result={**_result(), "output_file": str(output)},
    )
    docx = client.get(f"/api/bg/minutes/{task_id}", params={"format": "docx"})
    assert docx.status_code == 200
    assert "wordprocessingml" in docx.headers["content-type"]
    assert "filename*=UTF-8''" in docx.headers["content-disposition"]
    text = client.get(f"/api/bg/minutes/{task_id}")
    assert text.status_code == 200
    assert "会議の本文" in text.text
    rejected = client.get(f"/api/bg/minutes/{task_id}", params={"format": "pdf"})
    assert rejected.status_code == 400


def test_docx_contains_headings_table_and_speakers():
    content = build_minutes_docx(
        task_id="abc",
        name="定例",
        created_at="2026-10-03T00:00:00",
        result={
            **_result(),
            "segments": [
                {"start": 65, "end": 70, "text": "了解しました", "speaker": "山田"}
            ],
            "transcript": "[山田] 了解しました",
        },
    )
    document = Document(BytesIO(content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "定例" in text
    assert "要約" in text
    assert "議事録" in text
    assert "アクションアイテム" in text
    assert "文字起こし" in text
    assert "山田" in text
    assert "01:05" in text
    assert document.tables
    cells = [cell.text for cell in document.tables[0].rows[1].cells]
    assert "資料送付" in cells
