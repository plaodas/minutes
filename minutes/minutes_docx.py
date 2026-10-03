import io
import re
from datetime import datetime
from typing import Any

from docx import Document
from docx.shared import Pt

DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


def _text(value: object) -> str:
    return value.strip() if isinstance(value, str) else ""


def _timestamp(seconds: object) -> str:
    try:
        total = max(0, int(float(seconds)))
    except (TypeError, ValueError):
        return ""
    minutes, secs = divmod(total, 60)
    hours, minutes = divmod(minutes, 60)
    if hours:
        return f"{hours:02d}:{minutes:02d}:{secs:02d}"
    return f"{minutes:02d}:{secs:02d}"


def _add_markdownish(document: Document, text: str) -> None:
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("### "):
            document.add_heading(line[4:].strip(), level=2)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            document.add_heading(line[2:].strip(), level=2)
        elif line.startswith(("- ", "* ")):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            document.add_paragraph(line)


def safe_docx_filename(name: object, task_id: str) -> str:
    base = _text(name) or "minutes"
    base = re.sub(r'[\\/:*?"<>|\r\n]+', "_", base).strip(" .")[:80]
    if not base:
        base = "minutes"
    suffix = task_id.replace("-", "")[:8]
    return f"{base}-{suffix}.docx"


def build_minutes_docx(
    *,
    task_id: str,
    name: object,
    created_at: object,
    result: dict[str, Any],
) -> bytes:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = _text(name) or "議事録"
    document.add_heading(title, level=0)
    when = created_at
    if isinstance(when, datetime):
        when = when.isoformat(sep=" ", timespec="minutes")
    meta = _text(when) or task_id
    document.add_paragraph(meta)

    summary = _text(result.get("summary"))
    document.add_heading("要約", level=1)
    document.add_paragraph(summary or "要約はありません。")

    document.add_heading("議事録", level=1)
    minutes = _text(result.get("minutes"))
    if minutes:
        _add_markdownish(document, minutes)
    else:
        document.add_paragraph("議事録本文はありません。")

    document.add_heading("アクションアイテム", level=1)
    items = result.get("action_items")
    rows = items if isinstance(items, list) else []
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = table.rows[0].cells
    for cell, label in zip(
        headers, ("誰が", "何を", "いつまでに", "内容"), strict=False
    ):
        cell.text = label
    wrote_row = False
    for item in rows:
        if isinstance(item, str):
            values = ("", "", "", item)
        elif isinstance(item, dict):
            values = (
                _text(item.get("who")),
                _text(item.get("what")),
                _text(item.get("due") or item.get("when")),
                _text(item.get("text")),
            )
            if not any(values):
                continue
        else:
            continue
        cells = table.add_row().cells
        for cell, value in zip(cells, values, strict=False):
            cell.text = value
        wrote_row = True
    if not wrote_row:
        cells = table.add_row().cells
        cells[3].text = "アクションアイテムはありません。"

    document.add_heading("文字起こし", level=1)
    segments = result.get("segments")
    if isinstance(segments, list) and segments:
        for segment in segments:
            if not isinstance(segment, dict):
                continue
            speaker = _text(segment.get("speaker")) or "話者未設定"
            start = _timestamp(segment.get("start"))
            end = _timestamp(segment.get("end"))
            span = f"{start}–{end}" if start or end else ""
            heading = " ".join(part for part in (speaker, span) if part)
            document.add_paragraph(heading, style="List Bullet")
            body = _text(segment.get("text"))
            if body:
                document.add_paragraph(body)
    else:
        transcript = _text(result.get("transcript"))
        document.add_paragraph(transcript or "文字起こしはありません。")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
